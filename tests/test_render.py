import os
import subprocess
import sys
import time
from pathlib import Path

import pytest

from meltify.converters import quicklook, render


@pytest.fixture(autouse=True)
def no_quicklook(monkeypatch):
    # Tests that want Quick Look turn it back on, so a Mac runs the same as Linux CI
    monkeypatch.setattr(quicklook, "available", lambda: False)


def _pdf(path, width=200, height=100):
    import pymupdf

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    doc.new_page(width=width, height=height)
    doc.save(path)
    return path


def test_nothing_runs_without_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr(render.shutil, "which", lambda name: None)
    monkeypatch.setattr(render, "MAC_SOFFICE", tmp_path / "absent")
    monkeypatch.setenv("CLAUDE_PLUGIN_DATA", str(tmp_path / "data"))
    monkeypatch.setattr(render, "run", lambda *a, **k: pytest.fail("soffice ran"))
    assert not render.available()
    assert render.to_pdf(tmp_path / "a.emf") is None
    assert render.to_pdfs([tmp_path / "a.emf"], tmp_path / "out") == {}


def test_one_soffice_start_converts_every_file_without_a_shell(tmp_path, monkeypatch):
    calls = []

    def run(args, *, timeout=None, check=True):
        calls.append((args, timeout))
        out = tmp_path / "out"
        _pdf(out / "a.pdf")
        return subprocess.CompletedProcess(args, 0, "", "")

    monkeypatch.setattr(render, "soffice", lambda: "/opt/soffice")
    monkeypatch.setattr(render, "run", run)
    a, b = tmp_path / "a.emf", tmp_path / "b.wmf"
    got = render.to_pdfs([a, b], tmp_path / "out", timeout=7)
    # b.wmf wrote no PDF, so it's left out for the caller to report
    assert got == {a: tmp_path / "out" / "a.pdf"}
    [(args, timeout)] = calls
    assert timeout == 7
    assert args[0] == "/opt/soffice" and "--headless" in args
    assert args[-2:] == [str(a), str(b)]
    assert args[args.index("--outdir") + 1] == str(tmp_path / "out")
    profile = next(x for x in args if x.startswith("-env:UserInstallation="))
    assert profile.endswith("/out/.profile")


def test_files_named_like_options_reach_soffice_as_absolute_paths(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(render, "soffice", lambda: "/opt/soffice")
    monkeypatch.setattr(render, "run", lambda args, **kw: calls.append(args))
    monkeypatch.chdir(tmp_path)
    hostile = Path("--accept=socket,host=0.0.0.0,port=2002;urp;.wpd")
    hostile.write_bytes(b"x")
    render.converted([hostile], "pdf", tmp_path / "out")
    [args] = calls
    # Everything after the program is either a known option or an absolute path
    assert args[-1] == str(tmp_path.resolve() / hostile)
    assert not any(a.startswith("--accept") for a in args)
    assert args[args.index("--outdir") + 1] == str(tmp_path.resolve() / "out")


@pytest.mark.skipif(sys.platform == "win32", reason="process groups are POSIX")
def test_a_timed_out_run_kills_what_the_program_started(tmp_path):
    # Like soffice, a launcher that forks the real worker and waits on it
    pid_file = tmp_path / "worker.pid"
    launcher = (
        "import subprocess, sys, time\n"
        "w = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])\n"
        f"open({str(pid_file)!r}, 'w').write(str(w.pid))\n"
        "time.sleep(60)\n"
    )
    with pytest.raises(subprocess.TimeoutExpired):
        render.run([sys.executable, "-c", launcher], timeout=2)
    worker = int(pid_file.read_text())
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline and _alive(worker):
        time.sleep(0.05)
    assert not _alive(worker)


def _alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    return True


def test_managed_dir_is_removed_when_soffice_fails(tmp_path, monkeypatch):
    made = []
    real = render.tempfile.mkdtemp

    def mkdtemp(**kw):
        made.append(real(dir=tmp_path, **kw))
        return made[-1]

    def run(args, **kw):
        raise subprocess.TimeoutExpired(args, 1)

    monkeypatch.setattr(render, "soffice", lambda: "/opt/soffice")
    monkeypatch.setattr(render, "run", run)
    monkeypatch.setattr(render.tempfile, "mkdtemp", mkdtemp)
    with pytest.raises(subprocess.TimeoutExpired):
        render.to_pdf(tmp_path / "a.docx")
    assert made and not any(Path(p).exists() for p in made)


def test_missing_output_raises_but_keeps_a_caller_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(render, "soffice", lambda: "/opt/soffice")
    monkeypatch.setattr(render, "run", lambda args, **kw: None)
    with pytest.raises(RuntimeError, match=r"no \.pdf"):
        render.to_pdf(tmp_path / "a.docx", out_dir=tmp_path / "keep")
    assert (tmp_path / "keep").is_dir()


def test_png_crops_a_region_by_page_fractions(tmp_path):
    from io import BytesIO

    from PIL import Image

    pdf = _pdf(tmp_path / "a.pdf", 720, 540)
    whole = Image.open(BytesIO(render.png(pdf, dpi=72)))
    part = Image.open(BytesIO(render.png(pdf, 1, (0.5, 0.5, 1.0, 1.0), dpi=72)))
    assert whole.size == (720, 540)
    assert part.size == (360, 270)


@pytest.mark.skipif(render.soffice() is None, reason="LibreOffice is not installed")
def test_real_soffice_renders_a_docx(tmp_path):
    docx = pytest.importorskip("docx")

    path = tmp_path / "a.docx"
    doc = docx.Document()
    doc.add_paragraph("rendered by LibreOffice")
    doc.save(path)
    pdf = render.to_pdf(path, out_dir=tmp_path / "out")
    assert pdf is not None and pdf.read_bytes().startswith(b"%PDF")


def test_a_downloaded_soffice_comes_first(tmp_path, monkeypatch):
    monkeypatch.setenv("CLAUDE_PLUGIN_DATA", str(tmp_path))
    binary = tmp_path / "tools/libreoffice/program/soffice"
    binary.parent.mkdir(parents=True)
    binary.write_text("#!/bin/sh\n")
    binary.chmod(0o755)
    monkeypatch.setattr(render.shutil, "which", lambda name: "/usr/bin/soffice")
    assert render.soffice() == str(binary)
    assert render.renderers() == ["soffice"]


def _fake_quicklook(monkeypatch, calls):
    def to_pdf(path, out_dir=None, timeout=quicklook.TIMEOUT):
        calls.append(path)
        return _pdf(out_dir / f"{path.stem}.pdf")

    monkeypatch.setattr(quicklook, "available", lambda: True)
    monkeypatch.setattr(quicklook, "to_pdf", to_pdf)


def test_quicklook_draws_what_soffice_missed(tmp_path, monkeypatch):
    calls = []
    _fake_quicklook(monkeypatch, calls)

    def run(args, **kw):
        _pdf(tmp_path / "out" / "a.pdf")

    monkeypatch.setattr(render, "soffice", lambda: "/opt/soffice")
    monkeypatch.setattr(render, "run", run)
    a, b = tmp_path / "a.key", tmp_path / "b.pages"
    got = render.to_pdfs([a, b], tmp_path / "out")
    assert calls == [b]
    assert {p: render.MADE_BY[f][0] for p, f in got.items()} == {a: "soffice", b: "quicklook"}


def test_quicklook_stands_in_for_a_failed_soffice(tmp_path, monkeypatch):
    calls = []
    _fake_quicklook(monkeypatch, calls)

    def run(args, **kw):
        raise subprocess.TimeoutExpired(args, 1)

    monkeypatch.setattr(render, "soffice", lambda: "/opt/soffice")
    monkeypatch.setattr(render, "run", run)
    pdf = render.to_pdf(tmp_path / "a.key", out_dir=tmp_path / "out")
    assert pdf == tmp_path / "out" / "a.pdf" and render.MADE_BY[pdf][0] == "quicklook"
    # Nothing drew the file, so the soffice error is what the caller reports
    monkeypatch.setattr(quicklook, "to_pdf", lambda path, out_dir=None, timeout=0: None)
    with pytest.raises(subprocess.TimeoutExpired):
        render.to_pdf(tmp_path / "b.key", out_dir=tmp_path / "out")
    with pytest.raises(subprocess.TimeoutExpired):
        render.to_pdfs([tmp_path / "b.key"], tmp_path / "out")


def test_quicklook_alone_is_a_renderer(tmp_path, monkeypatch):
    calls = []
    _fake_quicklook(monkeypatch, calls)
    monkeypatch.setattr(render, "soffice", lambda: None)
    monkeypatch.setattr(render, "run", lambda *a, **k: pytest.fail("soffice ran"))
    assert render.available() and render.renderers() == ["quicklook"]
    pdf = render.to_pdf(tmp_path / "a.key")
    assert pdf is not None and pdf.is_file()
    import shutil

    shutil.rmtree(pdf.parent)


def test_quicklook_failure_without_soffice_is_none_and_cleans_up(tmp_path, monkeypatch):
    made = []
    real = render.tempfile.mkdtemp

    def mkdtemp(**kw):
        made.append(real(dir=tmp_path, **kw))
        return made[-1]

    monkeypatch.setattr(render.tempfile, "mkdtemp", mkdtemp)
    monkeypatch.setattr(render, "soffice", lambda: None)
    monkeypatch.setattr(quicklook, "available", lambda: True)
    monkeypatch.setattr(quicklook, "to_pdf", lambda path, out_dir=None, timeout=0: None)
    assert render.to_pdf(tmp_path / "a.key") is None
    assert made and not any(Path(p).exists() for p in made)


def _wmf(path: Path) -> Path:
    """A placeable WMF header and an empty record list, enough for python-pptx to size it"""
    import struct

    # Placeable header: key, handle, bbox in units, 1440 units per inch, reserved, checksum
    words = struct.pack("<IHhhhhHI", 0x9AC6CDD7, 0, 0, 0, 2880, 1440, 1440, 0)
    checksum = 0
    for (word,) in struct.iter_unpack("<H", words[:20]):
        checksum ^= word
    header = words[:20] + struct.pack("<H", checksum)
    # Standard header, then the end-of-file record
    body = struct.pack("<HHHIHIH", 1, 9, 0x300, 12, 0, 3, 0) + struct.pack("<IH", 3, 0)
    path.write_bytes(header + body)
    return path


def test_lone_metafile_is_wrapped_in_a_slide_for_quicklook(tmp_path, monkeypatch):
    pytest.importorskip("pptx")
    seen = []

    def to_pdf(path, out_dir=None, timeout=quicklook.TIMEOUT):
        from pptx import Presentation

        deck = Presentation(str(path))
        seen.append((path.suffix, deck.slide_width, deck.slide_height, len(deck.slides)))
        return _pdf(out_dir / f"{path.stem}.pdf")

    monkeypatch.setattr(quicklook, "available", lambda: True)
    monkeypatch.setattr(quicklook, "to_pdf", to_pdf)
    monkeypatch.setattr(render, "soffice", lambda: None)
    pic = _wmf(tmp_path / "v0.wmf")
    got = render.to_pdfs([pic], tmp_path / "out")
    assert got == {pic: tmp_path / "out" / "v0.pdf"}
    # A 2 by 1 inch picture gets a slide its own size
    assert seen == [(".pptx", 2 * 914400, 914400, 1)]


def test_wrap_without_python_pptx_is_skipped(tmp_path, monkeypatch):
    import builtins

    real = builtins.__import__

    def blocked(name, *args, **kwargs):
        if name == "pptx":
            raise ImportError(name)
        return real(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", blocked)
    assert render.wrap(tmp_path / "a.emf", tmp_path) is None
