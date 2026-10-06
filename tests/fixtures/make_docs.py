"""A mixed folder of documents like the ones handed over with a task"""

from __future__ import annotations

import io
import struct
import zipfile
from email.message import EmailMessage
from email.utils import make_msgid
from pathlib import Path


def calendar_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Calendar"
    ws.append(["date", "event"])
    ws.append(["2025-07-21", "handover meeting"])
    ws["B7"] = "budget review"
    secret = wb.create_sheet("Old")
    secret.sheet_state = "hidden"
    secret.append(["archived"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mixed_folder(root: Path) -> Path:
    import pymupdf
    from PIL import Image

    root.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Page one says the deadline is Friday", fontsize=12)
    doc.new_page()
    doc.save(root / "report.pdf")

    (root / "calendar.xlsx").write_bytes(calendar_xlsx())
    (root / "notes.txt").write_text("first line\n둘째 줄 한글\nthird line\n", encoding="utf-8")
    (root / "legacy.txt").write_bytes("완료 보고".encode("cp949"))
    Image.new("RGB", (40, 20), "white").save(root / "scan.png")
    (root / "blob.bin").write_bytes(b"\x00\x01\x02binary")

    msg = EmailMessage()
    msg["From"] = "lead@example.com"
    msg["To"] = "new@example.com"
    msg["Subject"] = "handover"
    msg.set_content("Please check the attached calendar.")
    msg.add_attachment(
        calendar_xlsx(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="cal.xlsx",
    )
    (root / "mail").mkdir()
    (root / "mail" / "handover.eml").write_bytes(msg.as_bytes())
    (root / ".hidden.txt").write_text("skip me")
    return root


def nested_mail(path: Path, levels: int) -> Path:
    """A mail forwarded inside itself `levels` times, with a text attachment at the bottom"""
    msg = EmailMessage()
    msg["Subject"] = "level 0"
    msg.set_content("innermost body")
    msg.add_attachment(b"deep note", maintype="text", subtype="plain", filename="deep.txt")
    for n in range(1, levels + 1):
        outer = EmailMessage()
        outer["Subject"] = f"level {n}"
        outer.set_content(f"forward {n}")
        outer.add_attachment(msg)
        msg = outer
    path.write_bytes(msg.as_bytes())
    return path


def tricky_mail(path: Path, absolute: Path) -> Path:
    """Attachment names that escape the output folder or repeat inside one mail"""
    msg = EmailMessage()
    msg["Subject"] = "tricky"
    msg.set_content("see attachments")
    for name, body in [
        ("../../../../escape.txt", b"escape"),
        (str(absolute), b"absolute"),
        ("..", b"dots"),
        ("dup.txt", b"first copy"),
        ("dup.txt", b"second copy"),
    ]:
        msg.add_attachment(body, maintype="text", subtype="plain", filename=name)
    path.write_bytes(msg.as_bytes())
    return path


def card(text: str, size: tuple[int, int] = (300, 100), color: str = "white") -> bytes:
    """A PNG with one line of text, the kind of picture OCR should read"""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", size, color)
    ImageDraw.Draw(img).text((10, size[1] // 2 - 5), text, fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def embedded_pdf(path: Path, logo: bytes | None = None) -> Path:
    """Four pages: a picture to read, a repeated logo, an icon, a tiny image and a scan

    1. page 1 has text, the invoice card at (72,100)-(522,250) and the logo
    2. page 2 has text, the logo and a 100 px image drawn at 30 by 30 pt
    3. page 3 is a full-page scan with no text layer
    4. page 4 has text, the logo and a 24 px image drawn large
    """
    import pymupdf

    logo = logo or card("LOGO", (100, 100), "#ddd")
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "Quarterly report, see the invoice below", fontsize=12)
    p1.insert_image(pymupdf.Rect(72, 100, 522, 250), stream=card("INVOICE 2026-0917", (900, 300)))
    p1.insert_image(pymupdf.Rect(500, 700, 560, 760), stream=logo)
    p2 = doc.new_page()
    p2.insert_text((72, 72), "Second page with only small pictures", fontsize=12)
    p2.insert_image(pymupdf.Rect(500, 700, 560, 760), stream=logo)
    p2.insert_image(pymupdf.Rect(72, 100, 102, 130), stream=card("ICON", (100, 100)))
    p3 = doc.new_page()
    p3.insert_image(p3.rect, stream=card("SCANNED PAGE 98,765", (850, 1100)))
    p4 = doc.new_page()
    p4.insert_text((72, 72), "Fourth page with a tiny picture drawn large", fontsize=12)
    p4.insert_image(pymupdf.Rect(500, 700, 560, 760), stream=logo)
    p4.insert_image(pymupdf.Rect(72, 100, 272, 300), stream=card("x", (24, 24)))
    doc.save(path)
    return path


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
IMAGE_REL = f"{R}/image"


def _picture(rid: str) -> str:
    return (
        f'<w:r><w:drawing><a:graphic xmlns:a="{A}"><a:graphicData><a:blip r:embed="{rid}"/>'
        "</a:graphicData></a:graphic></w:drawing></w:r>"
    )


def embedded_docx(path: Path, image: bytes) -> Path:
    """Paragraph 3 holds a picture, paragraph 4 an EMF and a linked picture, 5 a tiny icon"""
    paras = [
        "<w:p><w:r><w:t>Intro paragraph</w:t></w:r></w:p>",
        "<w:p><w:r><w:t>Second paragraph</w:t></w:r></w:p>",
        f"<w:p>{_picture('rId1')}</w:p>",
        f"<w:p>{_picture('rId2')}{_picture('rId3')}</w:p>",
        f"<w:p>{_picture('rId4')}</w:p>",
    ]
    document = (
        f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{W}" xmlns:r="{R}">'
        f"<w:body>{''.join(paras)}</w:body></w:document>"
    )
    rels = (
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{PKG}">'
        f'<Relationship Id="rId1" Type="{IMAGE_REL}" Target="media/image1.png"/>'
        f'<Relationship Id="rId2" Type="{IMAGE_REL}" Target="media/image2.emf"/>'
        f'<Relationship Id="rId3" Type="{IMAGE_REL}" Target="https://example.com/a.png"'
        ' TargetMode="External"/>'
        f'<Relationship Id="rId4" Type="{IMAGE_REL}" Target="/word/media/icon.png"/>'
        "</Relationships>"
    )
    types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/'
        'vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Default Extension="emf" ContentType="image/x-emf"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
        'officedocument.wordprocessingml.document.main+xml"/></Types>'
    )
    root_rels = (
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{PKG}">'
        f'<Relationship Id="rId1" Type="{R}/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", types)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("word/document.xml", document)
        z.writestr("word/_rels/document.xml.rels", rels)
        z.writestr("word/media/image1.png", image)
        z.writestr("word/media/image2.emf", b"\x01\x00\x00\x00 not really emf")
        z.writestr("word/media/icon.png", card("i", (16, 16)))
    return path


def embedded_pptx(path: Path, image: bytes) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for n in range(1, 4):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Slide {n}"
        if n == 3:
            slide.shapes.add_picture(io.BytesIO(card("i", (20, 20))), Inches(9), Inches(0.2))
            slide.shapes.add_picture(io.BytesIO(image), Inches(1), Inches(2), Inches(6))
    deck.save(path)
    return path


def embedded_xlsx(path: Path, image: bytes) -> Path:
    """The picture's top left corner sits on Sales!D2, written with openpyxl's absolute targets"""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImage

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["region", "amount"])
    ws.append(["Seoul", 100])
    ws.add_image(XImage(io.BytesIO(image)), "D2")
    wb.save(path)
    return path


def inline_image_mail(path: Path, image: bytes, attached: bytes) -> Path:
    """An HTML body showing a cid image from multipart/related, plus a regular attachment"""
    msg = EmailMessage()
    msg["From"] = "a@example.com"
    msg["To"] = "b@example.com"
    msg["Subject"] = "invoice inline"
    cid = make_msgid(domain="example.com")
    msg.set_content("plain body: see the invoice")
    msg.add_alternative(f'<p>see the invoice:</p><img src="cid:{cid[1:-1]}">', subtype="html")
    msg.get_payload()[1].add_related(
        image, maintype="image", subtype="png", cid=cid, filename="inv.png"
    )
    msg.add_attachment(attached, maintype="image", subtype="png", filename="chart.png")
    path.write_bytes(bytes(msg))
    return path


def _art(kind: int, body: bytes, ver: int = 0, inst: int = 0) -> bytes:
    """One OfficeArt record, a container when ver is 0xF"""
    return struct.pack("<HHI", ver | inst << 4, kind, len(body)) + body


def _picture_entry(png: bytes) -> bytes:
    """An OfficeArtFBSE holding a PNG blip after its fixed 36 bytes"""
    blip = _art(0xF01E, bytes(16) + b"\xff" + png, inst=0x6E0)
    head = bytes([6, 6]) + bytes(18) + struct.pack("<III", len(blip), 1, 0) + bytes(4)
    return _art(0xF007, head + blip, ver=2, inst=6)


def _shape(spid: int, pib: int | None) -> bytes:
    fsp = _art(0xF00A, struct.pack("<II", spid, 0xA00), ver=2, inst=75)
    opt = b"" if pib is None else _art(0xF00B, struct.pack("<HI", 0x4104, pib), ver=3, inst=1)
    return _art(0xF004, fsp + opt, ver=0xF)


def word97(path: Path, text: str, inline: bytes, floating: bytes) -> Path:
    """A Word 97 binary whose \\x01 shows `inline` and whose \\x08 anchors `floating`

    The inline picture sits in the Data stream where the character's sprmCPicLocation
    points, the floating one in the drawing store, reached through its anchor's shape
    """
    from hwpx.hwp5.cfb import build_compound_file

    start, page = 1024, 3
    end = start + 2 * len(text)
    pic = start + 2 * text.index("\x01")
    # Character runs before, on and after the picture character, the middle one pointing
    # at offset 0 of the Data stream and marked special
    fkp = bytearray(512)
    struct.pack_into("<4I", fkp, 0, start, pic, pic + 2, end)
    fkp[16:19] = bytes([0, 240, 0])
    chpx = struct.pack("<HI", 0x6A03, 0) + struct.pack("<HB", 0x0855, 1)
    fkp[480 : 481 + len(chpx)] = bytes([len(chpx)]) + chpx
    fkp[511] = 3

    sp = _shape(1025, 1)
    picf = struct.pack("<IHH", 0, 0x44, 0x64).ljust(0x44, b"\0") + sp + _picture_entry(inline)
    data = struct.pack("<I", len(picf)) + picf[4:]

    store = _art(0xF001, _picture_entry(floating), ver=0xF, inst=1)
    drawing = _art(0xF003, _shape(1024, None) + _shape(2049, 1), ver=0xF)
    dgg = _art(0xF000, store, ver=0xF) + b"\0" + _art(0xF002, drawing, ver=0xF)
    clx = struct.pack("<II", 0, len(text)) + struct.pack("<HIH", 0, start, 0)
    clx = b"\x02" + struct.pack("<I", len(clx)) + clx
    chpx_plc = struct.pack("<III", start, end, page)
    # One anchor naming shape 2049, then its 26 byte Spa
    spa = struct.pack("<III", text.index("\x08"), len(text), 2049) + bytes(22)
    table = bytearray()
    pairs = [(0, 0)] * 93
    for index, part in ((33, clx), (12, chpx_plc), (40, spa), (50, dgg)):
        pairs[index] = (len(table), len(part))
        table += part

    fib = bytearray(struct.pack("<HH", 0xA5EC, 0xC1).ljust(32, b"\0"))
    struct.pack_into("<H", fib, 10, 0x0200)  # the table stream is 1Table
    struct.pack_into("<II", fib, 0x18, start, end)  # fcMin and fcMac, where text sits
    longs = bytearray(88)
    struct.pack_into("<i", longs, 12, len(text))
    fib += struct.pack("<H", 14) + bytes(28) + struct.pack("<H", 22) + longs
    fib += struct.pack("<H", 93) + b"".join(struct.pack("<II", *p) for p in pairs) + bytes(2)
    doc = bytes(fib).ljust(start, b"\0") + text.encode("utf-16-le")
    doc = doc.ljust(page * 512, b"\0") + bytes(fkp)
    streams = {"WordDocument": doc, "1Table": bytes(table), "Data": data}
    path.write_bytes(build_compound_file(streams.items()))
    return path
