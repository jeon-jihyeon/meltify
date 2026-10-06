"""Renders for what native parsing can't reach, like EMF pictures

LibreOffice goes first, from PATH, the macOS app or the portable copy `meltify doctor
--install libreoffice` downloads. Quick Look covers what it misses on macOS
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import threading
from collections.abc import Sequence
from pathlib import Path

from meltify import tools
from meltify.converters import quicklook
from meltify.converters.run import LOCK, workdir
from meltify.safe import run

# Where the macOS app keeps the binary when it isn't linked onto PATH
MAC_SOFFICE = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
# Where `meltify doctor --install libreoffice` puts it in the tools dir
PORTABLE_SOFFICE = ("LibreOffice.app/Contents/MacOS/soffice", "libreoffice/program/soffice")
TIMEOUT = 120
DPI = 200
# Pictures Quick Look only previews from inside a slide
WRAPPED = (".emf", ".wmf")
# PowerPoint's slide size limits in EMU, 1 to 56 inches
SLIDE_MIN, SLIDE_MAX = 914400, 51206400
# Which renderer wrote each PDF and the sha256 of the file it drew, for needs and OCR
# cache keys. Output paths are unique temp paths, so a plain dict under a lock is enough
MADE_BY: dict[Path, tuple[str, str | None]] = {}
_MADE_LOCK = threading.Lock()


def soffice() -> str | None:
    found = tools.find("soffice", *PORTABLE_SOFFICE) or shutil.which("libreoffice")
    if found:
        return found
    return str(MAC_SOFFICE) if MAC_SOFFICE.exists() else None


def renderers() -> list[str]:
    """Renderers this machine has, in the order they're tried"""
    found = ["soffice"] if soffice() is not None else []
    return found + (["quicklook"] if quicklook.available() else [])


def available() -> bool:
    return bool(renderers())


def drawn_from(pdf: Path) -> str | None:
    """The renderer and the source content behind a PDF this module wrote

    Renders embed the time they were made, so their own bytes never repeat. This names
    what the pages show instead, and stays the same from one run to the next
    """
    with _MADE_LOCK:
        by, source = MADE_BY.get(pdf, (None, None))
    return None if source is None else f"{by}:{source}"


def _made(pdf: Path, by: str, source: Path) -> Path:
    from meltify.files import sha256

    # An iWork package folder has no single content to hash, so its pages key on the PDF
    content = sha256(source) if source.is_file() else None
    with _MADE_LOCK:
        # Callers remove most renders themselves, so a long-lived process would otherwise
        # keep every path it ever made
        for gone in [p for p in MADE_BY if not p.exists()]:
            del MADE_BY[gone]
        MADE_BY[pdf] = (by, content)
    return pdf


def converted(
    paths: Sequence[Path], target: str, out_dir: Path, timeout: float = TIMEOUT
) -> dict[Path, Path]:
    """Files soffice converted to `target`, like pdf or xlsx, from one start, keyed by input

    Inputs need distinct stems, since each output is named after its input. An input
    soffice couldn't open is left out, and a failed or timed out run raises
    """
    binary = soffice()
    if binary is None or not paths:
        return {}
    out_dir.mkdir(parents=True, exist_ok=True)
    # Absolute paths, since soffice takes a file named like `--accept=...` as an option, and
    # it has no `--` to end them
    outdir = out_dir.resolve()
    # A private profile per call, since parallel soffice runs fight over the shared one
    profile = (outdir / ".profile").as_uri()
    run(
        [
            binary,
            "--headless",
            "--norestore",
            f"-env:UserInstallation={profile}",
            "--convert-to",
            target,
            "--outdir",
            str(outdir),
            *(str(p.resolve()) for p in paths),
        ],
        timeout=timeout,
    )
    # A target like `txt:Text (encoded):UTF8` names a filter after the extension
    ext = target.split(":")[0]
    found = {p: out_dir / f"{p.stem}.{ext}" for p in paths}
    return {p: f for p, f in found.items() if f.exists()}


def convert_to(path: Path, target: str, out_dir: Path, timeout: float = TIMEOUT) -> Path:
    """One file converted to `target`, raising when soffice wrote nothing"""
    got = converted([path], target, out_dir, timeout).get(path)
    if got is None:
        raise RuntimeError(f"soffice wrote no .{target.split(':')[0]} for {path.name}")
    return got


def soffice_pdf(path: Path, out_dir: Path, timeout: float = TIMEOUT) -> Path:
    """One file drawn to PDF by soffice, recorded so OCR of its pages keys on `path`"""
    return _made(convert_to(path, "pdf", out_dir, timeout), "soffice", path)


def wrap(picture: Path, out_dir: Path) -> Path | None:
    """A one-slide pptx showing an EMF or WMF picture, or None without python-pptx

    Quick Look has no preview for a bare EMF or WMF, but draws one placed on a slide
    """
    try:
        from pptx import Presentation
    except ImportError:
        return None
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    try:
        pic = slide.shapes.add_picture(str(picture), 0, 0)
    except Exception:  # noqa: BLE001
        # python-pptx reads the header for the size, and a broken one fails in many ways
        return None
    width, height = int(pic.width), int(pic.height)
    if width <= 0 or height <= 0:
        return None
    # Fit the slide to the picture, scaled into the sizes PowerPoint allows
    scale = min(SLIDE_MAX / max(width, height), max(1.0, SLIDE_MIN / min(width, height)))
    width, height = (min(SLIDE_MAX, max(SLIDE_MIN, round(v * scale))) for v in (width, height))
    pic.width, pic.height = width, height
    deck.slide_width, deck.slide_height = width, height
    out_dir.mkdir(parents=True, exist_ok=True)
    deck.save(out_dir / f"{picture.stem}.pptx")
    return out_dir / f"{picture.stem}.pptx"


def preview(path: Path, out_dir: Path, timeout: float = TIMEOUT) -> Path | None:
    """PDF of a Quick Look preview, or None when Quick Look can't draw the file"""
    if path.suffix.lower() not in WRAPPED:
        pdf = quicklook.to_pdf(path, out_dir=out_dir, timeout=timeout)
        return None if pdf is None else _made(pdf, "quicklook", path)
    with tempfile.TemporaryDirectory(prefix="meltify-wrap-") as tmp:
        deck = wrap(path, Path(tmp))
        pdf = None if deck is None else quicklook.to_pdf(deck, out_dir=out_dir, timeout=timeout)
    return None if pdf is None else _made(pdf, "quicklook", path)


def to_pdfs(paths: Sequence[Path], out_dir: Path, timeout: float = TIMEOUT) -> dict[Path, Path]:
    """PDFs keyed by input, from soffice in one start and Quick Look for what it missed

    An input no renderer could draw is left out. Raises only when soffice failed and
    nothing else drew a single file, so the caller can report why
    """
    done: dict[Path, Path] = {}
    error: Exception | None = None
    if soffice() is not None:
        try:
            done = {
                p: _made(f, "soffice", p)
                for p, f in converted(paths, "pdf", out_dir, timeout).items()
            }
        except (OSError, RuntimeError, subprocess.SubprocessError) as e:
            error = e
    if quicklook.available():
        for p in paths:
            if p not in done and (pdf := preview(p, out_dir, timeout)) is not None:
                done[p] = pdf
    if error is not None and not done:
        raise error
    return done


def to_pdf(path: Path, *, out_dir: Path | None = None, timeout: float = TIMEOUT) -> Path | None:
    """PDF of one file, or None when no renderer is installed or Quick Look can't draw it

    A soffice failure raises unless Quick Look draws the file instead. Without `out_dir`
    the PDF lands in a fresh temp dir that the caller removes
    """
    if not available():
        return None
    managed = out_dir is None
    folder = workdir("meltify-render-") if out_dir is None else out_dir
    try:
        pdf = None
        error: Exception | None = None
        if soffice() is not None:
            try:
                pdf = soffice_pdf(path, folder, timeout)
            except (OSError, RuntimeError, subprocess.SubprocessError) as e:
                error = e
        if pdf is None and quicklook.available():
            pdf = preview(path, folder, timeout)
        if pdf is None and error is not None:
            raise error
    except BaseException:
        if managed:
            shutil.rmtree(folder, ignore_errors=True)
        raise
    if pdf is None and managed:
        shutil.rmtree(folder, ignore_errors=True)
    return pdf


def png(
    pdf: Path,
    page: int = 1,
    clip: tuple[float, float, float, float] | None = None,
    dpi: int = DPI,
) -> bytes:
    """One page as PNG, optionally cut to a region given as fractions of the page"""
    import pymupdf

    with LOCK, pymupdf.open(pdf) as doc:
        p = doc[page - 1]
        rect = None
        if clip is not None:
            r = p.rect
            x0, y0, x1, y1 = clip
            rect = pymupdf.Rect(
                r.x0 + x0 * r.width, r.y0 + y0 * r.height, r.x0 + x1 * r.width, r.y0 + y1 * r.height
            )
        return p.get_pixmap(dpi=dpi, clip=rect).tobytes("png")
