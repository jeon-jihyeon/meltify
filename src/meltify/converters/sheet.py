from __future__ import annotations

import contextlib
import functools
import importlib.util
import io
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable, Iterable, Iterator
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path
from typing import IO, Any

from meltify.converters import Block, Converted
from meltify.converters.embeds import Embeds
from meltify.converters.ooxml import HANCOM, NS, R_ID, hancom_need, objects, package, rels
from meltify.converters.ooxml_charts import Resolve, read_object, take
from meltify.converters.tables import escape_cell, lettered
from meltify.converters.xmlsafe import zip_xml
from meltify.evidence import Src
from meltify.safe import MissingTool

ANCHORS = ("twoCellAnchor", "oneCellAnchor", "absoluteAnchor")
# The type attribute of a cell holding an error value, in either quote
ERROR_TYPE = (b' t="e"', b" t='e'")
Row = tuple[Any, ...]
# Rows of one sheet by title, None for a title the workbook doesn't have
Rows = Callable[[str], list[Row] | None]
# A workbook on disk, or the bytes of one embedded in another document
Source = Path | bytes


def _cell(v: Any) -> str:
    # Plain text and numbers come first, since a sheet holds millions of them
    if isinstance(v, str):
        return escape_cell(v) if "|" in v or "\n" in v else v
    if isinstance(v, int | float):
        return str(v)
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="minutes").removesuffix(" 00:00")
    if isinstance(v, date):
        return v.isoformat()
    return escape_cell(str(v))


def value(v: Any) -> Any:
    # calamine reads every number as a float and an empty cell as "", while openpyxl keeps
    # ints as ints and empty cells as None. Past 1e16 a float prints in exponent form, the
    # way writers store such numbers and openpyxl reads them back
    if v.__class__ is float:
        return int(v) if v.is_integer() and -1e16 < v < 1e16 else v
    return None if v == "" else v


def _shows_errors(source: Source) -> bool:
    """Whether a worksheet of an Excel package holds an error value like #N/A

    calamine reads those as blank cells, so such a workbook goes to openpyxl instead. The
    sheets stream through in chunks, with a tail kept so an attribute split between two
    still matches. Only a cell's type is ever "e" in worksheet XML, and text holding the
    same characters is escaped, so a plain byte search is enough
    """
    with _opened(source) as f:
        if not zipfile.is_zipfile(f):
            return False
        with zipfile.ZipFile(f) as z:
            for info in z.infolist():
                if "/worksheets/" not in f"/{info.filename}" or not info.filename.endswith(".xml"):
                    continue
                with z.open(info) as member:
                    tail = b""
                    while chunk := member.read(1 << 20):
                        if any(e in tail + chunk for e in ERROR_TYPE):
                            return True
                        tail = chunk[-8:]
    return False


def _anchor_cell(anchor: Any) -> str | None:
    from openpyxl.utils import get_column_letter

    start = anchor.find("xdr:from", NS)
    if start is None:
        return None
    col, row = start.findtext("xdr:col", None, NS), start.findtext("xdr:row", None, NS)
    if col is None or row is None:
        return None
    return f"{get_column_letter(int(col) + 1)}{int(row) + 1}"


def _opened(source: Source) -> contextlib.AbstractContextManager[IO[bytes]]:
    # A file object, since openpyxl refuses names it doesn't know like .cell
    return io.BytesIO(source) if isinstance(source, bytes) else source.open("rb")


def sheets(source: Source) -> Iterator[tuple[str, bool, Iterable[Row]]]:
    """Each worksheet's title, whether it's hidden and its saved values from row 1

    calamine reads cells about ten times faster than openpyxl, which stays the reader when
    the office extra isn't installed or a sheet shows error values. Chart sheets hold no
    cells, so they're left out
    """
    if importlib.util.find_spec("python_calamine") is None or _shows_errors(source):
        yield from _openpyxl_sheets(source)
        return
    from python_calamine import CalamineWorkbook, SheetTypeEnum, SheetVisibleEnum

    with _opened(source) as f:
        wb = CalamineWorkbook.from_filelike(f)
    try:
        for meta in wb.sheets_metadata:
            if meta.typ != SheetTypeEnum.WorkSheet:
                continue
            found = wb.get_sheet_by_name(meta.name)
            hidden = meta.visible != SheetVisibleEnum.Visible
            if found.start is None:
                # An empty sheet, whose rows calamine panics on instead of yielding none
                yield meta.name, hidden, ()
                continue
            # Rows come from row 1 but start at the first used column, so the empty columns
            # before it go back in to keep column letters right
            left = (None,) * found.start[1]
            yield meta.name, hidden, (left + tuple(map(value, r)) for r in found.iter_rows())
    finally:
        wb.close()


def _openpyxl_sheets(source: Source) -> Iterator[tuple[str, bool, Iterable[Row]]]:
    from openpyxl import load_workbook

    with _opened(source) as f:
        wb = load_workbook(f, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                # Read-only mode yields placeholders without coordinates, so count from row 1
                rows = ws.iter_rows(min_row=1, values_only=True)
                yield ws.title, ws.sheet_state != "visible", rows
        finally:
            wb.close()


def table(rows: Iterable[Row], at: Src) -> Block | None:
    """A sheet's cells as a markdown table, None when it holds no value"""
    kept = (
        (n, list(map(_cell, r)))
        for n, r in enumerate(rows, start=1)
        if any(v is not None for v in r)
    )
    text = lettered(kept)
    return None if text is None else Block(at, text)


def rows_by_title(source: Source) -> Rows:
    """Rows of one sheet, read the first time a chart formula names it and kept after

    Only charts that saved no values ask, so most workbooks never hold a sheet's rows
    past its table. A workbook nothing can open answers None, which leaves its charts
    listed as having no cached values
    """

    @functools.cache
    def rows(title: str) -> list[Row] | None:
        try:
            for name, _, found in sheets(source):
                if name == title:
                    return list(found)
        except (KeyboardInterrupt, SystemExit):
            raise
        except BaseException:  # noqa: BLE001
            # A broken workbook neither reader opens. calamine reports some by panicking,
            # which isn't an Exception
            return None
        return None

    return rows


def recalculated(path: Path, src: Src) -> Embeds | None:
    """Charts read again after LibreOffice calculates the workbook, None without it

    openpyxl and other generators save formulas with no results, so neither the cells nor
    a chart built on them have a value until something calculates them. LibreOffice also
    caches every chart's values when it saves, which covers defined names and other
    series formulas the resolver can't follow
    """
    from meltify.converters import render, run

    # --shallow starts no subprocess, so the charts stay listed as having no cached values.
    # Only LibreOffice calculates, so another page renderer doesn't help here
    if run.current().shallow or render.soffice() is None:
        return None
    with tempfile.TemporaryDirectory(prefix="meltify-recalc-") as tmp:
        try:
            saved = render.convert_to(path, "xlsx", Path(tmp))
            embeds = images(path, src, resolver(rows_by_title(saved)))
            if embeds.uncached:
                _cached_in(saved, src, embeds)
        except (OSError, RuntimeError, subprocess.SubprocessError, zipfile.BadZipFile):
            return None
    return embeds


def _cached_in(saved: Path, src: Src, embeds: Embeds) -> None:
    """Fill uncached charts from the copy LibreOffice saved, matched by sheet and anchor cell

    The saved copy numbers its drawings its own way, so only where a chart sits ties it to
    the original
    """
    found: dict[tuple[str | None, str | None], list[str]] = {}
    try:
        again = images(saved, src)
    except (ET.ParseError, ValueError, KeyError):
        return
    for b in again.blocks:
        if b.text.startswith("Chart"):
            found.setdefault((b.src.sheet, b.src.cell), []).append(b.text)
    for at in list(embeds.uncached):
        if texts := found.get((at.sheet, at.cell)):
            embeds.uncached.remove(at)
            embeds.blocks.append(Block(at, texts.pop(0)))


def resolver(rows: Rows) -> Resolve:
    """Cell values a chart formula like `'Sales 2026'!$B$2:$B$5` names

    openpyxl writes charts with formulas and no cached values, so these fill them in
    """
    from openpyxl.utils.cell import range_to_tuple

    def resolve(ref: str) -> list[str] | None:
        try:
            title, (c0, r0, c1, r1) = range_to_tuple(ref.replace("$", ""))
        except (ValueError, TypeError):
            return None
        found = rows(title)
        if found is None or c0 is None or c1 is None:
            return None
        picked = found[(r0 or 1) - 1 : r1 or len(found)]
        return [_cell(r[c - 1] if c <= len(r) else None) for r in picked for c in range(c0, c1 + 1)]

    return resolve


def images(path: Path, src: Src, resolve: Resolve | None = None) -> Embeds:
    """Pictures, charts and diagrams in each sheet's drawing, cited by their top-left cell"""
    embeds = Embeds()
    with zipfile.ZipFile(path) as z:
        found = package(z)
        book = found.main if found else "xl/workbook.xml"
        parts = rels(z, book)
        sheets = zip_xml(z, book).find("s:sheets", NS)
        for sheet in [] if sheets is None else sheets:
            part = parts.get(sheet.get(R_ID) or "")
            if part is None or part not in z.namelist():
                continue
            name = sheet.get("name") or ""
            drawings = [
                t
                for t in rels(z, part).values()
                if t and t.startswith("xl/drawings/") and t.endswith(".xml")
            ]
            n = 0
            for drawing in drawings:
                if drawing not in z.namelist():
                    continue
                root = zip_xml(z, drawing)
                targets = rels(z, drawing)
                for anchor in root:
                    if anchor.tag.rsplit("}", 1)[-1] not in ANCHORS:
                        continue
                    at = replace(src, sheet=name, cell=_anchor_cell(anchor))
                    for kind, rid, _ in objects(anchor):
                        n += 1
                        take(embeds, z, kind, rid, targets, replace(at, img=n), resolve)
    return embeds


def _tables(path: Path, src: Src, kind: str) -> Converted:
    """Each worksheet's cells as a table, with hidden sheets listed in needs"""
    out = Converted(kind)
    hidden = []
    for title, is_hidden, rows in sheets(path):
        if is_hidden:
            hidden.append(title)
        if (block := table(rows, replace(src, sheet=title))) is not None:
            out.blocks.append(block)
    if hidden:
        out.needs.append("hidden sheets " + ", ".join(hidden))
    return out


def _xlsb_images(path: Path, src: Src, resolve: Resolve) -> Embeds:
    """Pictures and charts in an xlsb package, numbered across the workbook

    Sheet names and drawing anchors sit in binary records calamine doesn't expose, so they
    cite the workbook instead of a cell. Charts are DrawingML XML as in xlsx, and only the
    sheets are binary
    """
    embeds = Embeds()
    with zipfile.ZipFile(path) as z:
        n = 0
        for info in z.infolist():
            if info.filename.startswith("xl/charts/chart") and info.filename.endswith(".xml"):
                n += 1
                read_object(embeds, z, "chart", info.filename, replace(src, img=n), resolve)
            elif not info.filename.startswith("xl/media/"):
                continue
            elif (data := embeds.load(z, info.filename)) is not None:
                n += 1
                embeds.picture(replace(src, img=n), data, info.filename)
    return embeds


def convert_binary(path: Path, src: Src) -> Converted:
    """Excel 97 to 2003 workbooks and their templates, and Excel 2007 binary workbooks

    calamine reads both the same way, and the xlsb package also holds pictures and charts
    """
    if importlib.util.find_spec("python_calamine") is None:
        # openpyxl reads only the zip formats, so the binary ones need calamine
        raise MissingTool("python-calamine", "meltify doctor --install office")
    out = _tables(path, src, "legacy")
    if path.suffix.lower() not in (".xls", ".xlt"):
        _xlsb_images(path, src, resolver(rows_by_title(path))).into(out)
    return out


def convert(path: Path, src: Src) -> Converted:
    suffix = path.suffix.lower()
    if suffix in HANCOM and not zipfile.is_zipfile(path):
        return Converted("sheet", needs=[hancom_need(suffix)])
    out = _tables(path, src, "sheet")
    # Sheets are read again only for charts whose formulas name them
    embeds = images(path, src, resolver(rows_by_title(path)))
    if embeds.uncached and (calculated := recalculated(path, src)) is not None:
        # Charts are the only thing read again, so the cell tables keep what the file saved
        embeds = calculated
    embeds.into(out)
    return out
