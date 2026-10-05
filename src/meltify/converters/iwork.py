from __future__ import annotations

import importlib.util
import warnings
import zipfile
from dataclasses import replace
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter

from meltify.converters import Block, Child, Converted, RecognizeJob
from meltify.evidence import Src

# Newer iWork files keep a full-size preview at the root, older ones under QuickLook
PREVIEWS = ("preview.jpg", "QuickLook/Preview.jpg", "QuickLook/Thumbnail.jpg")
PREVIEW_PDF = "QuickLook/Preview.pdf"
MAX_PREVIEW = 256 * 1024 * 1024


def _cell(cell: Any) -> str:
    from numbers_parser import MergedCell

    # A merged range shows its value once, at the top-left cell
    if isinstance(cell, MergedCell):
        return ""
    return (cell.formatted_value or "").replace("|", "\\|").replace("\n", " ")


def _numbers(path: Path, src: Src, out: Converted) -> None:
    from numbers_parser import Document

    # numbers-parser warns about every document version it hasn't seen, which isn't actionable
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        doc = Document(path)
    for sheet in doc.sheets:
        for table in sheet.tables:
            rows = [
                (n, [_cell(c) for c in row]) for n, row in enumerate(table.iter_rows(), start=1)
            ]
            rows = [(n, r) for n, r in rows if any(r)]
            if not rows:
                continue
            width = max(len(r) for _, r in rows)
            letters = [get_column_letter(i + 1) for i in range(width)]
            # Row numbers and column letters let a reader cite a single cell
            lines = ["| row | " + " | ".join(letters) + " |", "|---" * (width + 1) + "|"]
            lines += [f"| {n} | " + " | ".join(r + [""] * (width - len(r))) + " |" for n, r in rows]
            table_src = replace(src, sheet=f"{sheet.name}>{table.name}")
            out.blocks.append(Block(table_src, "\n".join(lines)))


def _preview(path: Path, src: Src, out: Converted) -> bool:
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        for name in (*PREVIEWS, PREVIEW_PDF):
            if name not in names or z.getinfo(name).file_size > MAX_PREVIEW:
                continue
            data = z.read(name)
            if name == PREVIEW_PDF:
                out.children.append(Child(name, src, data))
            else:
                out.jobs.append(RecognizeJob("image", src.inside(name), data=data))
            return True
    return False


def convert(path: Path, src: Src) -> Converted:
    out = Converted("iwork")
    if path.suffix.lower() == ".numbers":
        if importlib.util.find_spec("numbers_parser") is None:
            out.needs.append("iwork extra")
        else:
            from numbers_parser import UnsupportedError

            try:
                _numbers(path, src, out)
                return out
            except UnsupportedError as e:
                if "encrypted" not in str(e):
                    raise
                out.needs.append("iwork encrypted")
                return out
    # Pages and Keynote text lives in undocumented protobuf, so OCR the preview image
    if _preview(path, src, out):
        out.needs.append("iwork preview only")
    else:
        out.needs.append("iwork preview missing")
    return out
